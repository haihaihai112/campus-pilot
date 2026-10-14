from __future__ import annotations

from pathlib import Path


SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg"}


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
    for table in document.tables:
        paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(paragraphs)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    values: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                values.append(f"第{slide_number}页：{shape.text.strip()}")
    return "\n".join(values)


def _image(path: Path) -> str:
    try:
        from paddleocr import PaddleOCR
    except ImportError as exc:
        raise RuntimeError("Image OCR requires optional dependency: pip install paddleocr") from exc
    result = PaddleOCR(use_doc_orientation_classify=True, use_doc_unwarping=True).predict(str(path))
    texts: list[str] = []
    for item in result:
        data = getattr(item, "json", item)
        if isinstance(data, dict):
            texts.extend(data.get("rec_texts", []))
    return "\n".join(texts)


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8-sig", errors="ignore")
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return _image(path)
    raise ValueError(f"Unsupported document type: {path.suffix}")


def discover(directory: Path) -> list[Path]:
    paths = [
        path
        for path in directory.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED and path.name.lower() != "readme.md"
    ]
    if not paths:
        raise FileNotFoundError(f"No supported knowledge files found in {directory}")
    return sorted(paths)

